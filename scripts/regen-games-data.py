#!/usr/bin/env python3
"""Regenerate assets/js/games-data.js from selected CSV + xlsx (Top 20)."""
import csv, json, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'])
    import openpyxl

ROOT = Path(__file__).resolve().parents[2]  # cps-games
SITE = Path(__file__).resolve().parents[1]
CSV = ROOT / 'pc-web-promo-games-selected.csv'
XLSX = ROOT / '2026-10-21_微信小游戏CPS游戏库.xlsx'

with CSV.open(encoding='utf-8-sig') as f:
    selected = list(csv.DictReader(f))[:20]

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ws = wb['全部游戏']
header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
col = {h: i for i, h in enumerate(header)}
wanted = {r['appid'] for r in selected}
found = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    aid = row[col['appid']]
    if aid in wanted:
        found[aid] = row
        if len(found) == len(wanted):
            break

games = []
for r in selected:
    aid = r['appid']
    row = found.get(aid)
    def cell(name, default=''):
        if not row:
            return default
        v = row[col[name]]
        return default if v is None else v
    def share(v):
        try:
            return float(v) if v not in (None, '') else None
        except Exception:
            return None
    games.append({
        'appid': aid,
        'name': r.get('appname') or cell('appname'),
        'intro': r.get('brief_intro') or cell('brief_intro') or '',
        'logo': cell('logo URL') or '',
        'banner': cell('banner URL') or '',
        'category': f"{r.get('二级分类名') or cell('二级分类名')}/{r.get('三级分类名') or cell('三级分类名')}".strip('/'),
        'shareNew': share(r.get('新用户分成%') or cell('新用户分成%')),
        'shareReturn': share(r.get('回流用户分成%') or cell('回流用户分成%')),
    })

out = SITE / 'assets/js/games-data.js'
out.write_text(
    '/* Auto-generated Top 20. */\nwindow.GAMES_DATA = '
    + json.dumps(games, ensure_ascii=False, indent=2)
    + ';\n',
    encoding='utf-8',
)
print('wrote', out, 'n=', len(games))
